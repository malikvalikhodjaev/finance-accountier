"""Read Payme's verified 16-column XLSX export without changing it."""
import re
import zipfile
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

HEADERS = ['Дата платежа', 'Время платежа', 'Тип операции', 'Имя поставщика',
           'Название организации поставщика', 'Сумма платежа', 'Категория',
           'Получатель / Отправитель', 'Реквизиты платежа', 'Номер карты',
           'Название карты', 'Комментарий к платежу', 'Тип чека', 'Состояние чека',
           'Пункт обслуживания', 'Терминал']


def extract(rows, numeric, filename, sheet='Filtered_Cheques'):
    if not rows or list(rows[0]) != HEADERS or len(rows) > 10001:
        raise ValueError('Не распознан Excel Payme: нужны 16 колонок штатной выгрузки.')
    period_match = re.fullmatch(r'(\d{8})_(\d{8})(?: \(\d+\))?\.xlsx', Path(filename).name, re.I)
    period = None
    if period_match:
        start, end = [datetime.strptime(value, '%Y%m%d').strftime('%Y-%m-%d') for value in period_match.groups()]
        if start > end:
            raise ValueError('Некорректный период в имени файла Payme.')
        period = {'from': start, 'to': end}
    result = []
    occurrences = Counter()
    for index, values in enumerate(rows[1:], 2):
        values = list(values)
        if len(values) != 16:
            raise ValueError('Неожиданное число колонок Excel Payme.')
        if not any(value is not None and value != '' for value in values):
            continue
        if any(isinstance(value, str) and value.startswith('=') for value in values):
            raise ValueError('Вместо исходной выгрузки получен файл с формулами. Выбери оригинал Payme.')
        if any(value is not None and not isinstance(value, str) for i, value in enumerate(values) if i != 5):
            raise ValueError('Неожиданный тип значения в строке Excel Payme.')
        fields = [value or '' for value in values]
        paid = fields[13] == 'Оплачен'
        if fields[13] not in ('Оплачен', 'Отменен') or fields[2] not in ('Списание', 'Поступление') or fields[12] != 'Payme':
            raise ValueError('Неизвестный тип или состояние чека Payme. Нужна проверка формата.')
        if not fields[3] or not re.fullmatch(r'(?:\d{6}\*{6}\d{4}|\*{16})', fields[9]):
            raise ValueError('Не распознаны поставщик или маска карты Payme.')
        try:
            amount = Decimal(numeric[index])
        except (InvalidOperation, KeyError):
            raise ValueError('Сумма Payme должна быть исходным числом Excel.') from None
        minor = amount * 100
        if not minor.is_finite() or minor != minor.to_integral_value() or not 0 <= minor <= 100000000000000 or (paid and minor == 0):
            raise ValueError('Не распознана точная сумма Payme с двумя десятичными знаками.')
        day = datetime.strptime(fields[0], '%d-%m-%Y').strftime('%Y-%m-%d') if fields[0] else None
        time = fields[1] or None
        if time:
            datetime.strptime(time, '%H:%M:%S')
        if paid and (not day or not time):
            raise ValueError('У оплаченного чека отсутствуют дата или время.')
        if period and day and not period['from'] <= day <= period['to']:
            raise ValueError('Дата чека выходит за период в имени выгрузки Payme.')
        # There is no cheque ID in this export. Keep repeated identical rows as
        # separate occurrences, and match only unchanged evidence on later import.
        identity = [day, time, fields[2], fields[3].strip(), fields[4].strip(), int(minor),
                    fields[8].strip(), fields[9], fields[12], fields[13], fields[14].strip(), fields[15].strip()]
        import json
        key = json.dumps(identity, ensure_ascii=False, separators=(',', ':'))
        occurrences[key] += 1
        result.append({'documentDate': day, 'documentNumber': str(index), 'postingDate': day,
                       'postingTime': time, 'amountMinor': int(minor), 'currency': 'UZS',
                       'details': fields[2] + ': ' + fields[3].strip(), 'operation': fields[2],
                       'merchant': fields[3].strip(), 'category': fields[6].strip(),
                       'cardSuffix': fields[9][-4:] if fields[9][-4:].isdigit() else None,
                       'paid': paid, 'identity': identity, 'occurrence': occurrences[key],
                       'fragments': [{'sheet': sheet, 'row': index, 'cells': values, 'amountExcel': numeric[index]}]})
    if not result:
        raise ValueError('В Excel Payme нет операций.')
    dates = [r['postingDate'] for r in result if r['postingDate']]
    if not dates:
        raise ValueError('В Excel Payme не найдены даты операций.')
    return {'source': 'payme-xlsx-v1', 'period': period or {'from': min(dates), 'to': max(dates)},
            'coverage': {'from': min(dates), 'to': max(dates)}, 'sheets': 1, 'currency': 'UZS', 'rows': result}


def read_xlsx(filename, original_name):
    import openpyxl
    with zipfile.ZipFile(filename) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(item.file_size for item in entries) > 32 * 1024 * 1024:
            raise ValueError('Слишком большой распакованный Excel. Выгрузи меньший период.')
        if 'xl/vbaProject.bin' in archive.namelist():
            raise ValueError('Нужен исходный XLSX без макросов.')
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        numeric = {}
        for cell in xml.findall('.//s:sheetData/s:row/s:c', ns):
            if re.fullmatch(r'F\d+', cell.attrib.get('r', '')) and cell.attrib.get('t', 'n') == 'n' and cell.find('s:f', ns) is None:
                value = cell.find('s:v', ns)
                if value is not None and value.text:
                    numeric[int(cell.attrib['r'][1:])] = value.text
    workbook = openpyxl.load_workbook(filename, read_only=True, data_only=False, keep_links=False)
    try:
        if workbook.sheetnames != ['Filtered_Cheques']:
            raise ValueError('Нужен исходный Excel Payme с листом Filtered_Cheques.')
        sheet = workbook['Filtered_Cheques']
        if sheet.max_row > 10001 or sheet.max_column != 16:
            raise ValueError('Неподдерживаемый размер таблицы Payme. Максимум 10 000 строк.')
        return extract(list(sheet.values), numeric, original_name)
    finally:
        workbook.close()
